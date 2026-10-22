"""Analyze the Money Manager Excel file structure"""
import openpyxl
import os
from collections import Counter

FILE = "/home/z/my-project/upload/Money Manager_22-10-2026.xlsx"

print(f"📂 File: {FILE}")
print(f"📊 Size: {os.path.getsize(FILE):,} bytes")
print("---")

wb = openpyxl.load_workbook(FILE, data_only=True)
print(f"📑 Sheets: {wb.sheetnames}")
for ws_name in wb.sheetnames:
    ws = wb[ws_name]
    print(f"\n📋 Sheet '{ws_name}': {ws.max_row} rows × {ws.max_column} cols")

# Analyze first sheet (Money Manager)
ws = wb[wb.sheetnames[0]]
print("\n=== HEADERS ===")
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
for i, h in enumerate(headers):
    print(f"  Col {i+1}: {h}")

print("\n=== SAMPLE ROWS (first 3) ===")
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=4, values_only=True)):
    print(f"  Row {i+1}:")
    for j, val in enumerate(row):
        print(f"    {headers[j]}: {val}")
    print()

# Collect unique values for key columns
print("\n=== UNIQUE ACCOUNTS (col 2) ===")
accounts = Counter()
categories = Counter()
types = Counter()
currencies = Counter()
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[1]: accounts[row[1]] += 1
    if row[2]: categories[row[2]] += 1
    if row[6]: types[row[6]] += 1
    if row[10]: currencies[row[10]] += 1

print(f"\n🔢 {len(accounts)} unique accounts:")
for acc, count in accounts.most_common():
    print(f"  {acc}: {count}")

print(f"\n🔢 {len(categories)} unique categories:")
for cat, count in categories.most_common():
    print(f"  {cat}: {count}")

print(f"\n🔢 Types:")
for t, count in types.most_common():
    print(f"  {t}: {count}")

print(f"\n🔢 Currencies:")
for c, count in currencies.most_common():
    print(f"  {c}: {count}")

# Date format analysis
print("\n=== DATE FORMAT ===")
for row in ws.iter_rows(min_row=2, max_row=5, values_only=True):
    print(f"  Date: {row[0]} | type: {type(row[0]).__name__}")
